from datetime import date, datetime, time

import pytest
from openpyxl import Workbook

from check_timesheets.main import _read_timesheet
from check_timesheets.timesheet_rates import TimesheetRates

EVENTS = [("GALA Day Rate", 47.0), ("GALA Half Day Rate", 28.0), ("House Event", 39.0)]


def _put(ws, row, col, value):
    ws.cell(row=row, column=col + 1, value=value)


def _legacy_workbook(path):
    wb = Workbook()
    ws = wb.active
    _put(ws, 3, 1, "Employee Name:"); _put(ws, 3, 2, "Sam Swimmer"); _put(ws, 3, 4, "Payroll Month:"); _put(ws, 3, 5, "September")
    for i, h in enumerate(["Date", "Day of", "Start", "End", "House", "Level"]):
        _put(ws, 5, i, h)
    for i, v in enumerate([datetime(2026, 9, 9), "Wed", time(19), time(20), "Acton", "L2 (Enhanced)"]):
        _put(ws, 6, i, v)
    _put(ws, 20, 6, 0.03)  # rate increase
    _put(ws, 21, 2, "Standard rates of pay (exclusive of holiday pay) ")
    levels = [("L2", 19.65), ("L2 (Enhanced)", 22.44), ("Admin", 12.21)]
    for i, (lvl, prev) in enumerate(levels):
        _put(ws, 22 + i, 4, lvl); _put(ws, 22 + i, 5, 0); _put(ws, 22 + i, 6, prev)
    for i, (lvl, rate) in enumerate(EVENTS):
        _put(ws, 22 + i, 9, lvl); _put(ws, 22 + i, 10, rate)
    wb.save(path)


def _current_workbook(path, rows):
    wb = Workbook()
    ws = wb.active
    _put(ws, 3, 1, "Employee Name*:"); _put(ws, 3, 2, "Sam Swimmer"); _put(ws, 3, 4, "Payroll Month*:"); _put(ws, 3, 5, "October")
    for i, h in enumerate(["Date*", "Day of", "Start", "End", "House*", "Level*"]):
        _put(ws, 5, i, h)
    for r, (d, level) in enumerate(rows):
        for i, v in enumerate([d, "Wed", time(19), time(20), "Acton", level]):
            _put(ws, 6 + r, i, v)
    _put(ws, 25, 4, "TOTALS FOR PRE OCTOBER 2026")
    _put(ws, 30, 1, "Standard rates of pay (before holiday pay) ")
    _put(ws, 30, 4, "2026/27"); _put(ws, 30, 5, "2025/26"); _put(ws, 30, 6, "2024/25")
    levels = [("Lifeguard", 11.07, 10.85), ("L2", 20.85, 20.24), ("L2 (Enhanced)", 23.81, 23.11), ("Admin", 13.08, 12.82)]
    for i, (lvl, new, old) in enumerate(levels):
        _put(ws, 31 + i, 3, lvl); _put(ws, 31 + i, 4, new); _put(ws, 31 + i, 5, old)
    _put(ws, 31 + len(levels), 3, "Other")
    for i, (lvl, rate) in enumerate(EVENTS):
        _put(ws, 31 + i, 9, lvl); _put(ws, 31 + i, 10, rate)
    wb.save(path)


def test_legacy_template_uses_one_rate_with_percentage_increase(tmp_path):
    path = tmp_path / "old.xlsx"
    _legacy_workbook(path)
    name, entries, month = _read_timesheet(str(path))
    assert name == "Sam Swimmer"
    assert month == "September"
    (entry,) = entries
    assert entry.rate == round(22.44 * 1.03, 2)
    assert entry.hours == 1.0


def test_current_template_switches_rate_on_change_date(tmp_path):
    path = tmp_path / "new.xlsx"
    _current_workbook(path, [(datetime(2026, 9, 30), "L2"), (datetime(2026, 10, 1), "L2")])
    name, entries, month = _read_timesheet(str(path))
    assert month == "October"
    rates = {e.date: e.rate for e in entries}
    assert rates == {date(2026, 9, 30): 20.24, date(2026, 10, 1): 20.85}


def test_current_template_unknown_level_lists_valid_levels(tmp_path):
    path = tmp_path / "new.xlsx"
    _current_workbook(path, [(datetime(2026, 10, 2), "L9")])
    with pytest.raises(ValueError, match="invalid level 'L9'.*admin"):
        _read_timesheet(str(path))


def test_unknown_template_is_reported(tmp_path):
    wb = Workbook(); ws = wb.active
    _put(ws, 3, 2, "Sam")
    for i, h in enumerate(["Date", "Day of", "Start", "End", "House", "Level"]):
        _put(ws, 5, i, h)
    path = tmp_path / "x.xlsx"; wb.save(path)
    with pytest.raises(ValueError, match="rate table header"):
        _read_timesheet(str(path))


def test_timesheet_rates_without_change_ignores_date():
    rates = TimesheetRates({"l2": 20.0})
    assert rates.rate_for("l2", date(2030, 1, 1)) == 20.0
    assert rates.rate_for("nope", date(2030, 1, 1)) is None
